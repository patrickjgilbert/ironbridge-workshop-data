#!/usr/bin/env python3
"""QA for the Ironbridge workshop pack. Rebuilds the model, reads the generated
files back the way an attendee would (CSV, xlsx, PDF text, docx) and asserts
every planted fact, every tie-out and the forbidden-term scan. Exits non-zero
on any failure and prints every planted number as it runs.

Run from the repo root:
    uv run --with openpyxl --with fpdf2 --with python-docx python3 scripts/qa_check.py
"""
import csv
import os
import re
import subprocess
import sys
import zipfile
from collections import defaultdict

from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate as g  # noqa: E402

ROOT = g.ROOT
D01 = g.D01
FAILS = []
PASSES = 0


def check(cond, msg):
    global PASSES
    if cond:
        PASSES += 1
        print("  ok   " + msg)
    else:
        FAILS.append(msg)
        print("  FAIL " + msg)


def pdftotext(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def read_csv(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


print("QA: Ironbridge workshop pack")
model = g.build_model()

# ------------------------------------------------------------ files present
print("\n[files]")
names = sorted(os.listdir(D01))
activity = [fn for _, fn in g.ACTIVITY_FILES]
check(all(fn in names for fn in activity), "13 activity files present with the exact names")
expected = set(activity) | {"merchant master.xlsx", "Pricing Schedule 2026.pdf", "Risk Policy memo.docx", g.BLANK_SCAN,
                            "RE_ return rate notice.txt", "FW_ pricing amendment - Harrow Finch.txt", "Northgate - processor review.txt",
                            "lunch order thursday.txt", "teams-chat-merchant-ops-sept.txt"}
messy = [n for n in names if n not in g.KEEP]
check(set(messy) == expected and len(messy) == 22, "exactly 22 messy files plus README.md and brand-guidelines.md (%d messy)" % len(messy))
check(all(os.path.exists(os.path.join(D01, k)) for k in g.KEEP), "README.md and brand-guidelines.md exist in 01-portfolio-inbox")
a, b = [open(os.path.join(D01, fn), "rb").read() for fn in g.DUPLICATE_FILE]
check(a == b and len(a) > 1000, "PLANTED  '%s' is byte-identical to '%s'" % (g.DUPLICATE_FILE[1], g.DUPLICATE_FILE[0]))

# ------------------------------------------------------------ activity exports, read like an attendee
print("\n[activity exports]")
data = {}   # (month, merchant_id) -> row
by_month = defaultdict(list)
seen_files = set()
for i, fn in g.ACTIVITY_FILES:
    if fn == g.DUPLICATE_FILE[1]:
        continue
    rows = read_csv(os.path.join(D01, fn))
    with open(os.path.join(D01, fn)) as f:
        header = f.readline().strip()
    check(header == ",".join(g.HEADER), "%s: exact header row" % fn)
    label = g.month_label(i)
    check(all(r["month"] == label for r in rows), "%s: every row is month %s" % (fn, label))
    want = 24 if i >= 10 else 23
    check(len(rows) == want, "%s: %d rows (one per merchant active that month)" % (fn, want))
    for r in rows:
        data[(i, r["merchant_id"])] = r
        by_month[i].append(r)
    seen_files.add(fn)

ints = ("debit_count", "credit_count", "returns_total", "returns_unauthorized", "returns_administrative", "returns_nsf")
check(all(re.fullmatch(r"\d+", r[k]) for r in data.values() for k in ints), "every count column is an integer")
check(all(re.fullmatch(r"\d+\.\d\d", r[k]) for r in data.values() for k in ("debit_volume", "credit_volume", "card_volume")), "every volume column has two decimals")
ratio_ok = True
for r in data.values():
    three = int(r["returns_unauthorized"]) + int(r["returns_administrative"]) + int(r["returns_nsf"])
    tot = int(r["returns_total"])
    if not (tot > three and (tot / three <= 1.10 if three < 200 else 1.015 <= tot / three <= 1.065)):
        ratio_ok = False
check(ratio_ok, "returns_total exceeds the three coded buckets in every row, by 2 to 6 percent wherever the counts are large enough to measure it")

# ------------------------------------------------------------ merchant master
print("\n[merchant master]")
wb = load_workbook(os.path.join(D01, "merchant master.xlsx"))
check(wb.sheetnames == ["Merchants"], "one sheet named Merchants")
ws = wb["Merchants"]
hdr = [c.value for c in ws[1]]
check(hdr == ["merchant_id", "merchant_name", "vertical", "onboarded", "account_manager", "pricing_tier", "contract_renewal", "status"], "master columns in order")
master = [dict(zip(hdr, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
check(len(master) == 24, "24 merchants in the master")
verts = defaultdict(int)
for m in master:
    verts[m["vertical"]] += 1
check(verts == {"Consumer lending": 7, "Auto finance": 3, "Collections / ARM": 6, "BNPL": 3, "E-commerce": 4, "Property management": 1},
      "verticals: 7 consumer lending, 3 auto, 6 ARM, 3 BNPL, 4 e-commerce, 1 property management")
check(len({m["account_manager"] for m in master}) == 4, "four account managers")
statuses = defaultdict(int)
for m in master:
    statuses[m["status"]] += 1
check(statuses == {"Active": 22, "Onboarding": 1, "Notice given": 1}, "status: 22 Active, 1 Onboarding, 1 Notice given")
mm = {m["merchant_id"]: m for m in master}
check(mm[g.HARROW]["pricing_tier"] == "B", "PLANTED  Harrow & Finch is still Tier B in the master")
check(mm[g.CEDAR]["status"] == "Notice given", "PLANTED  Cedar Falls Collections status is 'Notice given'")
check(mm[g.REDFERN]["status"] == "Onboarding" and str(mm[g.REDFERN]["onboarded"])[:7] == "2026-08", "Redfern Apparel is Onboarding, joined 2026-08")
check(mm[g.NORTHGATE]["pricing_tier"] == "A" and sum(1 for m in master if m["pricing_tier"] == "A") == 1, "Northgate is the only Tier A merchant")
names_in_files = {r["merchant_name"] for r in data.values()}
check(names_in_files == {m["merchant_name"] for m in master}, "merchant names in the exports match the master exactly")

# ------------------------------------------------------------ planted numbers, computed from the files
print("\n[planted numbers]")
tiers = {m["merchant_id"]: m["pricing_tier"] for m in master}
ttm_vol = sum(float(r["debit_volume"]) for r in data.values())
ttm_rev = 0.0
per_vol = defaultdict(float)
per_rev = defaultdict(float)
for (i, mid), r in data.items():
    fee, bps = g.TIERS[tiers[mid]]
    rev = int(r["debit_count"]) * fee + float(r["debit_volume"]) * bps / 10000 + int(r["credit_count"]) * g.CREDIT_FEE + float(r["card_volume"]) * g.CARD_BPS / 10000
    ttm_rev += rev
    per_vol[mid] += float(r["debit_volume"])
    per_rev[mid] += rev
may = sum(float(r["debit_volume"]) for r in by_month[7])
print("  TTM debit volume $%s (May once); May twice adds %.1f%%" % (f"{ttm_vol:,.2f}", 100 * may / ttm_vol))
print("  TTM revenue at master tiers $%s" % f"{ttm_rev:,.2f}")
check(2.0e9 <= ttm_vol <= 2.2e9, "TTM debit volume is about $2.1B ($%s)" % f"{ttm_vol:,.0f}")
check(abs(ttm_vol - model["ttm_vol"]) < 1, "TTM debit volume matches the model to the dollar")
check(2.8e6 <= ttm_rev <= 3.0e6, "TTM revenue at the tiers in the master is $2.8M to $3.0M ($%s)" % f"{ttm_rev:,.0f}")
check(abs(ttm_rev - model["ttm_rev"]) < 1, "TTM revenue matches the model to the dollar")
check(7.5 <= 100 * may / ttm_vol <= 9.0, "PLANTED  counting May twice inflates TTM volume by about 8%% (%.1f%%)" % (100 * may / ttm_vol))

# 1. Summit Ridge
sr = {i: data[(i, g.SUMMIT)] for i in range(12)}
rates = {i: round(100.0 * int(sr[i]["returns_unauthorized"]) / int(sr[i]["debit_count"]), 2) for i in range(12)}
print("  Summit Ridge unauthorized: " + "  ".join("%s %.2f%%" % (g.month_label(i), rates[i]) for i in range(12)))
for i, want in g.SUMMIT_RATES.items():
    check(rates[i] == want, "PLANTED  Summit Ridge unauthorized rate %s = %.2f%%" % (g.month_name(i), want))
    u, d = int(sr[i]["returns_unauthorized"]), int(sr[i]["debit_count"])
    check(abs(100.0 * u / d - want) < 0.004, "Summit Ridge %s rate sits away from a rounding boundary (%d / %d)" % (g.month_name(i), u, d))
check(all(rates[i] < 0.40 for i in range(8)), "Summit Ridge was under 0.40%% Oct through May (max %.2f%%)" % max(rates[i] for i in range(8)))
adm = {i: 100.0 * int(sr[i]["returns_administrative"]) / int(sr[i]["debit_count"]) for i in range(12)}
ovr = {i: 100.0 * int(sr[i]["returns_total"]) / int(sr[i]["debit_count"]) for i in range(12)}
check(max(adm.values()) < 2.5 and max(ovr.values()) < 12.0, "Summit Ridge administrative (max %.2f%%) and overall (max %.2f%%) stay under the internal lines" % (max(adm.values()), max(ovr.values())))
others_over = [(g.month_label(i), mid) for (i, mid), r in data.items() if mid != g.SUMMIT and 100.0 * int(r["returns_unauthorized"]) / int(r["debit_count"]) >= 0.40]
check(others_over == [], "every other merchant is under 0.40% unauthorized in every month")
other_rates = [100.0 * int(r["returns_unauthorized"]) / int(r["debit_count"]) for (i, mid), r in data.items() if mid != g.SUMMIT]
check(min(other_rates) >= 0.03 and max(other_rates) <= 0.33, "other merchants' unauthorized rates sit between 0.03%% and 0.33%% (%.2f to %.2f)" % (min(other_rates), max(other_rates)))
check(mm[g.SUMMIT]["pricing_tier"] == "B" and mm[g.SUMMIT]["vertical"] == "Consumer lending", "Summit Ridge is consumer lending, Tier B")

# 2. Northgate concentration
ranked = sorted(per_vol.items(), key=lambda kv: -kv[1])
ng_share = round(100.0 * per_vol[g.NORTHGATE] / ttm_vol, 1)
second_share = round(100.0 * ranked[1][1] / ttm_vol, 1)
print("  Northgate share %.1f%%, next largest %s %.1f%%, Northgate TTM revenue $%s" % (ng_share, mm[ranked[1][0]]["merchant_name"], second_share, f"{per_rev[g.NORTHGATE]:,.2f}"))
check(ng_share == 28.4, "PLANTED  Northgate BNPL is 28.4%% of TTM debit volume (%.1f%%)" % ng_share)
check(ranked[0][0] == g.NORTHGATE and 9.5 <= second_share <= 12.0, "next largest merchant is about 11%% (%.1f%%)" % second_share)
check(max(per_rev, key=per_rev.get) == g.NORTHGATE, "PLANTED  Northgate TTM revenue is the single largest line")
check(all(float(data[(i, g.NORTHGATE)]["debit_volume"]) > 25e6 for i in range(12)), "Northgate is over $25M every month (Tier A)")

# 3. Harrow & Finch under-billing
hf = sum(int(data[(i, g.HARROW)]["debit_count"]) * 0.03 + float(data[(i, g.HARROW)]["debit_volume"]) * 0.0004 for i in range(6, 12))
hf = round(hf, 2)
print("  Harrow & Finch under-billed Apr-Sep: $%s" % f"{hf:,.2f}")
check(hf == model["harrow_under"], "PLANTED  Harrow & Finch under-billing matches the generator to the cent ($%s)" % f"{hf:,.2f}")
check(38000 <= hf <= 45000, "Harrow & Finch under-billing is in the $38,000 to $45,000 band")
check(all(float(data[(i, g.HARROW)]["debit_volume"]) < 5e6 for i in range(6, 12)), "Harrow & Finch is under $5M every month Apr-Sep")
check(all(float(data[(i, g.HARROW)]["debit_volume"]) > 5e6 for i in range(0, 3)), "Harrow & Finch was over $5M in Oct-Dec 2025")

# 5. Cedar Falls
cq2 = sum(float(data[(i, g.CEDAR)]["debit_volume"]) for i in g.Q2)
cq3 = sum(float(data[(i, g.CEDAR)]["debit_volume"]) for i in g.Q3)
cpct = round(100.0 * (cq3 / cq2 - 1))
print("  Cedar Falls Q2 $%s -> Q3 $%s (%d%%)" % (f"{cq2:,.2f}", f"{cq3:,.2f}", cpct))
check(cpct == -38, "PLANTED  Cedar Falls Q3 debit volume is 38%% under Q2 (%d%%)" % cpct)
cv = [float(data[(i, g.CEDAR)]["debit_volume"]) for i in g.Q3]
check(cv[0] > cv[1] > cv[2], "Cedar Falls declines each month Jul, Aug, Sep")

# 6. administrative drift
def port_rate(i, col):
    return 100.0 * sum(int(r[col]) for r in by_month[i]) / sum(int(r["debit_count"]) for r in by_month[i])

a0, a11 = round(port_rate(0, "returns_administrative"), 1), round(port_rate(11, "returns_administrative"), 1)
print("  Portfolio administrative rate: " + "  ".join("%.2f" % port_rate(i, "returns_administrative") for i in range(12)))
check(a0 == 1.9 and a11 == 2.6, "PLANTED  portfolio administrative return rate %.1f%% (Oct 2025) -> %.1f%% (Sep 2026)" % (a0, a11))
check(all(port_rate(i, "returns_administrative") < 3.0 for i in range(12)), "portfolio administrative rate stays under Nacha's 3.0% all year")
check(all(100.0 * int(r["returns_administrative"]) / int(r["debit_count"]) < 3.0 for r in data.values()), "no merchant crosses Nacha's 3.0% administrative threshold in any month")
check(all(100.0 * int(r["returns_total"]) / int(r["debit_count"]) < 12.0 for r in data.values()), "no merchant crosses the 12% internal overall line in any month")
u_sep = port_rate(11, "returns_unauthorized")
check(u_sep < 0.40, "portfolio unauthorized rate in Sep 2026 is under 0.40%% (%.2f%%)" % u_sep)

# ------------------------------------------------------------ documents
print("\n[documents]")
ptxt = pdftotext(os.path.join(D01, "Pricing Schedule 2026.pdf"))
for need in ("$0.12", "8 bps", "$0.15", "11 bps", "$0.18", "15 bps", "$0.10", "22 basis points", "January 1, 2026", "Under $5,000,000"):
    check(need in ptxt, "pricing PDF text contains '%s'" % need)
btxt = pdftotext(os.path.join(D01, g.BLANK_SCAN)).strip()
check(btxt == "", "PLANTED  %s extracts no text (blank page)" % g.BLANK_SCAN)
doc = Document(os.path.join(D01, "Risk Policy memo.docx"))
dtxt = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
for need in ("0.40%", "0.5%", "2.5%", "3.0%", "12%", "15.0%", "25%", "Warning letter", "remediation", "Termination review", "R05, R07, R10, R11, R29", "R02, R03, R04", "R01"):
    check(need in dtxt, "risk policy memo contains '%s'" % need)

emails = {fn: open(os.path.join(D01, fn)).read() for fn in ("RE_ return rate notice.txt", "FW_ pricing amendment - Harrow Finch.txt", "Northgate - processor review.txt")}
for fn, txt in emails.items():
    head = txt.split("\n\n")[0].split("\n")
    check(head[0].startswith("From: ") and head[1].startswith("To: ") and any(l.startswith("Date: ") for l in head) and any(l.startswith("Subject: ") for l in head),
          "%s: From/To/Date/Subject header block" % fn)
e1 = emails["RE_ return rate notice.txt"]
check("September 24, 2026" in e1 and "June" in e1 and "authorize" in e1 and "entries" in e1, "Summit Ridge reply: dated Sep 24, disputes, asks for entries, mentions the June timing change")
check(("%.2f%%" % rates[10]) in e1, "Summit Ridge notice quotes the August rate that the data produces (%.2f%%)" % rates[10])
e2 = emails["FW_ pricing amendment - Harrow Finch.txt"]
check("March 27, 2026" in e2 and "Tier B to Tier C" in e2 and "April 1, 2026" in e2 and "$0.18" in e2 and "15 basis points" in e2, "Harrow & Finch amendment: Mar 27, Tier B to Tier C effective April 1 2026, Tier C pricing stated")
e3 = emails["Northgate - processor review.txt"]
check("September 18, 2026" in e3 and "November 15" in e3 and "two other providers" in e3 and "best and final" in e3 and "SLA" in e3, "Northgate email: Sep 18, two other providers, decision Nov 15, best and final, SLA")
ng_aug = float(data[(10, g.NORTHGATE)]["debit_volume"]) / 1e6
check(("$%d million" % int(ng_aug)) in e3 and ng_aug - int(ng_aug) < 0.75, "Northgate email quotes its August volume consistent with the data (a little over $%d million; actual %.2f)" % (int(ng_aug), ng_aug))
lunch = open(os.path.join(D01, "lunch order thursday.txt")).read()
check(re.search(r"20\d\d", lunch) is None and lunch.count("\n") >= 8, "PLANTED  lunch order has eight orders and no date")
chat = open(os.path.join(D01, "teams-chat-merchant-ops-sept.txt")).read()
for need in ("pulled May twice when the report timed out, one of them is just a copy, ignore it",
             "climbing again", "Nacha letter",
             "internal warning line is 0.40% unauthorized, tighter than Nacha's 0.50%",
             "new pricing ever make it into the master", "went quiet since July"):
    check(need in chat, "Teams export contains '%s'" % need[:60])
msgs = [l for l in chat.split("\n") if l.startswith("[2026-09-")]
check(25 <= len(msgs) <= 35, "Teams export has 25 to 35 messages in September 2026 (%d)" % len(msgs))
hf_q = [l for l in msgs if "make it into the master" in l]
check(hf_q and not any("yes" in l.lower() or "updated" in l.lower() for l in msgs[msgs.index(hf_q[0]) + 1: msgs.index(hf_q[0]) + 3]), "the Harrow & Finch question in the chat gets no answer")

# ------------------------------------------------------------ READMEs, skills, brand
print("\n[readmes and context]")
root_rd = open(os.path.join(ROOT, "README.md")).read()
f_rd = open(os.path.join(D01, "README.md")).read()
brand = open(os.path.join(D01, "brand-guidelines.md")).read()
prompts = ["This folder is the merchant portfolio inbox my analyst left me and it is a mess.",
           "Now build ironbridge-portfolio-review.xlsx from the activity exports",
           "Now turn that workbook into one self-contained HTML dashboard that I can open offline",
           "Turn what we just did in this session into a skill called merchant-portfolio-review",
           "Now reference brand-guidelines.md in this folder and follow those guidelines exactly as written.",
           "Search the web for Nacha's current ACH return rate thresholds"]
for p in prompts:
    check(p in root_rd and p in f_rd, "both READMEs carry the prompt '%s...'" % p[:50])
for need in ("0.31", "0.44", "0.52", "0.61", "28.4%", f"{hf:,.2f}", "38%", "1.9%", "2.6%", g.BLANK_SCAN, "lunch order thursday.txt", "Copy of activity_may_v2.csv",
             f"${ttm_vol:,.2f}", f"${ttm_rev:,.2f}", f"${per_rev[g.NORTHGATE]:,.2f}", f"${per_vol[g.NORTHGATE]:,.2f}", f"${cq2:,.2f}", f"${cq3:,.2f}"):
    check(need in f_rd, "folder README states the planted fact '%s'" % need)
for img in ("docs/img/01-code-button.png", "docs/img/02-download-zip.png", "docs/img/03-desktop-folder.png"):
    check(img in root_rd and os.path.exists(os.path.join(ROOT, img)), "root README shows %s and the file exists" % img)
hexes = set(re.findall(r"#[0-9A-Fa-f]{6}\b", brand))
check(len(hexes) >= 6, "brand-guidelines.md has hex color codes (%d)" % len(hexes))
check("fonts.googleapis.com" in brand and "pie or donut" in brand, "brand-guidelines.md has a Google Fonts link and the no-pie rule")
for sk in ("02-skills/daily-digest/SKILL.md", "02-skills/meeting-prep/SKILL.md"):
    p = os.path.join(ROOT, sk)
    check(os.path.exists(p) and open(p).read().startswith("---\nname:"), "%s exists with frontmatter" % sk)
text_files = [os.path.join(ROOT, "README.md"), os.path.join(D01, "README.md"), os.path.join(D01, "brand-guidelines.md"),
              os.path.join(ROOT, "02-skills/daily-digest/SKILL.md"), os.path.join(ROOT, "02-skills/meeting-prep/SKILL.md")] + \
             [os.path.join(D01, n) for n in messy if n.endswith(".txt")]
dashes = [os.path.relpath(p, ROOT) for p in text_files if os.path.exists(p) and ("—" in open(p).read() or "–" in open(p).read())]
check(dashes == [], "no em or en dashes in READMEs, skills, brand guidelines or text files (%s)" % (dashes or "clean"))
gi = open(os.path.join(ROOT, ".gitignore")).read()
for need in ("activity/", "reference/", "correspondence/", "policy/", "_review/", "ironbridge-portfolio-review.xlsx", "*.html"):
    check(need in gi, ".gitignore ignores attendee output %s" % need)

# ------------------------------------------------------------ hygiene
print("\n[hygiene]")
# stored reversed so a plain grep of the repo for these terms finds nothing, including this file
FORBIDDEN = [t[::-1] for t in ["ecnailyap", "ycrad", "nelluc", "dniwhtron", "anad"]]
hits = []
for dirpath, dirnames, filenames in os.walk(ROOT):
    dirnames[:] = [d for d in dirnames if d not in (".git", "__pycache__")]
    for fn in filenames:
        p = os.path.join(dirpath, fn)
        blobs = [open(p, "rb").read().lower()]
        if fn.lower().endswith(".pdf"):
            blobs.append(pdftotext(p).lower().encode())
        if fn.lower().endswith((".xlsx", ".docx")):
            with zipfile.ZipFile(p) as z:
                blobs += [z.read(n).lower() for n in z.namelist()]
        for term in FORBIDDEN:
            if any(re.search(rb"\b" + term.encode() + rb"\b", b) for b in blobs):
                hits.append((os.path.relpath(p, ROOT), term))
check(hits == [], "no forbidden terms anywhere in the repo (%s)" % (hits if hits else "clean"))
with zipfile.ZipFile(os.path.join(D01, "merchant master.xlsx")) as z:
    core = z.read("docProps/core.xml").decode()
    creators = set(a or b for a, b in re.findall(r"<dc:creator>(.*?)</dc:creator>|<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core))
    check(creators == {"Ironbridge"} and not any("comments" in n.lower() for n in z.namelist()), "merchant master.xlsx: author is Ironbridge, no comments")
check(doc.core_properties.author == "Ironbridge", "Risk Policy memo.docx: author is Ironbridge")
total = 0
for dirpath, dirnames, filenames in os.walk(ROOT):
    dirnames[:] = [d for d in dirnames if d not in (".git", "__pycache__")]
    total += sum(os.path.getsize(os.path.join(dirpath, f)) for f in filenames)
check(total < 1.5 * 1024 * 1024, "total size under 1.5 MB (%.2f MB)" % (total / 1024 / 1024))

print("\n%d checks passed, %d failed" % (PASSES, len(FAILS)))
for f in FAILS:
    print("  FAIL " + f)
sys.exit(1 if FAILS else 0)
