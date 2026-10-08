#!/usr/bin/env python3
"""Generate the Ironbridge workshop pack (folder 01-portfolio-inbox) from ONE
seeded model so every number ties out across the activity exports, the
merchant master, the pricing schedule and the correspondence.

Run from the repo root:
    uv run --with openpyxl --with fpdf2 --with python-docx python3 scripts/generate.py

Everything produced is fiction. Ironbridge Payments, every merchant, person
and dollar figure is invented for a training exercise.
"""
import csv
import os
import random
import shutil
import sys
from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content  # noqa: E402

SEED = 20261007
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D01 = os.path.join(ROOT, "01-portfolio-inbox")

# Twelve months, Oct 2025 -> Sep 2026. TTM = all twelve. Q2 = Apr-Jun, Q3 = Jul-Sep.
MONTHS = [(2025, 10), (2025, 11), (2025, 12)] + [(2026, m) for m in range(1, 10)]
Q2 = [6, 7, 8]
Q3 = [9, 10, 11]

# The thirteen activity files: month index -> filename. May appears twice.
ACTIVITY_FILES = [
    (0, "ach_activity_2025-10.csv"),
    (1, "ach_activity_2025-11.csv"),
    (2, "export (1).csv"),
    (3, "export (2).csv"),
    (4, "ach-activity-feb.csv"),
    (5, "ach activity march 2026.csv"),
    (6, "activity_april_FINAL.csv"),
    (7, "activity_may_v2.csv"),
    (7, "Copy of activity_may_v2.csv"),
    (8, "june.csv"),
    (9, "ach_activity_2026-07.csv"),
    (10, "aug export.csv"),
    (11, "Sept_export (1).csv"),
]
DUPLICATE_FILE = ("activity_may_v2.csv", "Copy of activity_may_v2.csv")
BLANK_SCAN = "scan_0097.pdf"
HEADER = ["month", "merchant_id", "merchant_name", "vertical", "debit_count", "debit_volume",
          "credit_count", "credit_volume", "returns_total", "returns_unauthorized",
          "returns_administrative", "returns_nsf", "card_volume"]

# Pricing schedule, effective Jan 1 2026. (item fee, bps on debit volume)
TIERS = {"A": (0.12, 8), "B": (0.15, 11), "C": (0.18, 15)}
CREDIT_FEE = 0.10
CARD_BPS = 22

CONSUMER, AUTO, ARM, BNPL, ECOM, PM = ("Consumer lending", "Auto finance", "Collections / ARM",
                                       "BNPL", "E-commerce", "Property management")
AMS = ["Priya Raman", "Tom Hallstrom", "Jenna Okoye", "Luis Ferreira"]

# Seasonality by vertical, Oct..Sep, multiplier on base monthly volume.
SEASON = {
    CONSUMER: [1.00, 0.98, 0.97, 1.03, 1.05, 1.06, 1.02, 1.00, 0.99, 1.00, 1.01, 1.00],
    AUTO:     [1.00, 0.99, 0.98, 1.00, 1.02, 1.04, 1.03, 1.02, 1.01, 1.01, 1.00, 0.99],
    ARM:      [0.97, 0.94, 0.92, 1.04, 1.12, 1.10, 1.03, 0.99, 0.96, 0.95, 0.96, 0.98],
    BNPL:     [1.00, 1.12, 1.22, 0.96, 0.90, 0.95, 0.98, 1.00, 1.01, 1.00, 1.03, 1.02],
    ECOM:     [0.98, 1.15, 1.25, 0.92, 0.88, 0.94, 0.97, 0.99, 1.00, 0.98, 1.02, 1.00],
    PM:       [1.00, 1.00, 1.00, 1.01, 1.01, 1.01, 1.01, 1.01, 1.02, 1.02, 1.02, 1.02],
}

# id, name, vertical, onboarded, account manager, tier in the master, renewal, status,
# base monthly debit volume ($M), average debit ticket ($), monthly growth, unauth %, admin %, nsf %,
# credit count ratio, credit ticket, card volume as a share of debit volume
MERCHANTS = [
    ("M-1001", "Northgate BNPL", BNPL, "2022-03-14", "Priya Raman", "A", "2026-12-31", "Active", None, 680, 0.012, 0.14, 2.45, 3.4, 0.02, 95, 0.006),
    ("M-1002", "Blue Heron Consumer Finance", CONSUMER, "2021-06-01", "Tom Hallstrom", "B", "2027-05-31", "Active", 17.64, 1300, 0.004, 0.17, 2.5, 6.4, 0.008, 310, 0.0025),
    ("M-1003", "Summit Ridge Lending", CONSUMER, "2023-09-18", "Tom Hallstrom", "B", "2027-09-30", "Active", 8.33, 540, 0.006, 0.21, 2.0, 6.8, 0.01, 240, 0.0025),
    ("M-1004", "Harrow & Finch Receivables", ARM, "2020-11-02", "Jenna Okoye", "B", "2027-03-31", "Active", None, 32, 0.0, 0.11, 2.7, 7.0, 0.003, 60, 0.0015),
    ("M-1005", "Cedar Falls Collections", ARM, "2022-08-22", "Jenna Okoye", "C", "2026-11-30", "Notice given", None, 430, 0.0, 0.13, 2.65, 6.9, 0.003, 110, 0.0015),
    ("M-1006", "Keystone Credit Partners", CONSUMER, "2021-02-15", "Luis Ferreira", "B", "2027-02-28", "Active", 12.74, 1340, 0.003, 0.19, 2.4, 6.1, 0.009, 290, 0.0025),
    ("M-1007", "Larkspur Lending", CONSUMER, "2024-04-08", "Tom Hallstrom", "C", "2027-04-30", "Active", 3.43, 1130, 0.009, 0.23, 2.6, 7.2, 0.011, 230, 0.003),
    ("M-1008", "Pinecrest Installment Loans", CONSUMER, "2022-10-03", "Luis Ferreira", "B", "2026-10-31", "Active", 6.86, 1060, 0.002, 0.16, 2.3, 6.6, 0.009, 260, 0.0025),
    ("M-1009", "Meridian Personal Finance", CONSUMER, "2023-01-23", "Priya Raman", "C", "2027-01-31", "Active", 2.65, 980, 0.005, 0.2, 2.55, 7.0, 0.01, 220, 0.003),
    ("M-1010", "Oakhaven Credit Services", CONSUMER, "2020-05-11", "Luis Ferreira", "C", "2027-06-30", "Active", 3.72, 1190, 0.001, 0.12, 2.3, 5.9, 0.007, 250, 0.002),
    ("M-1011", "Tri-County Auto Credit", AUTO, "2019-09-30", "Jenna Okoye", "B", "2027-08-31", "Active", 14.7, 940, 0.003, 0.09, 2.3, 4.6, 0.006, 410, 0.0015),
    ("M-1012", "Westbrook Motor Acceptance", AUTO, "2021-11-08", "Priya Raman", "B", "2026-12-31", "Active", 8.62, 890, 0.004, 0.11, 2.4, 4.9, 0.006, 380, 0.0015),
    ("M-1013", "Copperline Auto Finance", AUTO, "2023-05-15", "Tom Hallstrom", "C", "2027-05-31", "Active", 3.14, 860, 0.006, 0.14, 2.45, 5.1, 0.007, 360, 0.0015),
    ("M-1014", "Rivermark Recovery Group", ARM, "2020-02-03", "Jenna Okoye", "B", "2027-01-31", "Active", 6.17, 520, 0.002, 0.15, 2.6, 7.1, 0.003, 90, 0.0015),
    ("M-1015", "Stonebridge Receivables", ARM, "2022-01-10", "Luis Ferreira", "C", "2026-10-31", "Active", 2.35, 490, 0.001, 0.18, 2.7, 7.0, 0.003, 85, 0.0015),
    ("M-1016", "Northfield Collection Services", ARM, "2023-07-24", "Jenna Okoye", "C", "2027-07-31", "Active", 1.76, 470, 0.004, 0.22, 2.75, 6.8, 0.004, 80, 0.0015),
    ("M-1017", "Halvorsen & Dade Recovery", ARM, "2024-02-05", "Luis Ferreira", "C", "2027-02-28", "Active", 1.18, 430, 0.007, 0.25, 2.7, 6.9, 0.004, 75, 0.0015),
    ("M-1018", "Fernwood Pay", BNPL, "2024-06-17", "Priya Raman", "C", "2027-06-30", "Active", 3.53, 540, 0.01, 0.19, 2.3, 3.6, 0.022, 70, 0.01),
    ("M-1019", "Ashgrove Installments", BNPL, "2023-03-06", "Priya Raman", "B", "2027-03-31", "Active", 5.68, 600, 0.006, 0.16, 2.2, 3.2, 0.018, 80, 0.01),
    ("M-1020", "Lakeline Outdoor Supply", ECOM, "2022-04-25", "Tom Hallstrom", "C", "2027-04-30", "Active", 1.96, 400, 0.004, 0.08, 2.2, 1.6, 0.055, 95, 0.09),
    ("M-1021", "Brightwater Nutrition", ECOM, "2023-10-16", "Luis Ferreira", "C", "2026-10-31", "Active", 1.47, 310, 0.008, 0.1, 2.3, 1.8, 0.06, 60, 0.1),
    ("M-1022", "Juniper Home Goods", ECOM, "2024-09-09", "Tom Hallstrom", "C", "2027-09-30", "Active", 0.98, 370, 0.009, 0.07, 2.1, 1.5, 0.05, 90, 0.09),
    ("M-1023", "Redfern Apparel", ECOM, "2026-08-11", "Priya Raman", "C", "2028-08-31", "Onboarding", None, 350, 0.0, 0.09, 2.2, 1.6, 0.05, 80, 0.09),
    ("M-1024", "Elmhurst Property Management", PM, "2021-08-30", "Jenna Okoye", "B", "2027-08-31", "Active", 5.1, 2280, 0.002, 0.05, 1.8, 2.4, 0.004, 900, 0.0),
]
FIELDS = ["id", "name", "vertical", "onboarded", "am", "tier", "renewal", "status",
          "base", "ticket", "growth", "unauth", "admin", "nsf", "credit_ratio", "credit_ticket", "card_share"]

SUMMIT = "M-1003"
HARROW = "M-1004"
CEDAR = "M-1005"
NORTHGATE = "M-1001"
REDFERN = "M-1023"

SUMMIT_RATES = {8: 0.31, 9: 0.44, 10: 0.52, 11: 0.61}   # Jun, Jul, Aug, Sep 2026, percent
NORTHGATE_SHARE = 0.284
CEDAR_Q3_VS_Q2 = -0.38
ADMIN_DRIFT = (1.9, 2.6)   # portfolio administrative return rate, Oct 2025 -> Sep 2026, percent


def month_label(i):
    return "%04d-%02d" % MONTHS[i]


def month_name(i):
    return date(MONTHS[i][0], MONTHS[i][1], 1).strftime("%B %Y")


def cents(x):
    return round(x + 1e-9, 2)


def pct(num, den):
    return round(100.0 * num / den, 2)


def solve_rate(rng, rate, approx_count):
    """Find (debit_count, unauthorized) near approx_count whose rate, in percent
    rounded to 2 decimals, is exactly `rate`, and sits well away from a rounding
    boundary so Excel's ROUND and Python's round agree."""
    cands = []
    for d in range(approx_count - 400, approx_count + 400):
        u = int(round(rate * d / 100.0))
        if abs(100.0 * u / d - rate) < 0.003:
            cands.append((d, u))
    return rng.choice(cands)


def build_model():
    """The whole data model: per merchant per month rows, plus the planted figures.
    Deterministic for the seed, no file writes, so qa_check.py can rebuild it."""
    rng = random.Random(SEED)
    merchants = [dict(zip(FIELDS, m)) for m in MERCHANTS]
    byid = {m["id"]: m for m in merchants}
    rows = {}   # (month_index, merchant_id) -> dict

    def active(m, i):
        if m["id"] == REDFERN:
            return i >= 10
        return True

    # ---- debit volume and count, everyone except Northgate (set from the others' total)
    for m in merchants:
        for i in range(12):
            if not active(m, i):
                continue
            if m["id"] == HARROW:
                plan = [5.62, 5.41, 5.33, 5.14, 5.22, 4.97, 4.86, 4.92, 4.81, 4.78, 4.84, 4.69]
                vol = plan[i] * 1e6 * rng.uniform(0.985, 1.015)
            elif m["id"] == CEDAR:
                plan = [3.36, 3.22, 3.14, 3.52, 3.71, 3.63, 3.25, 3.30, 3.10, None, None, None]
                vol = plan[i] * 1e6 * rng.uniform(0.985, 1.015) if plan[i] else 0.0
            elif m["id"] == REDFERN:
                vol = [0.62, 1.31][i - 10] * 1e6 * rng.uniform(0.98, 1.02)
            elif m["id"] == NORTHGATE:
                vol = 0.0
            else:
                vol = m["base"] * 1e6 * SEASON[m["vertical"]][i] * ((1 + m["growth"]) ** i) * rng.uniform(0.975, 1.025)
            rows[(i, m["id"])] = {"vol": cents(vol)}

    # Cedar Falls: Q3 exactly 38% under Q2, falling each month
    q2 = sum(rows[(i, CEDAR)]["vol"] for i in Q2)
    q3 = cents(q2 * (1 + CEDAR_Q3_VS_Q2))
    jul = cents(q3 * 0.44)
    aug = cents(q3 * 0.33)
    rows[(9, CEDAR)]["vol"] = jul
    rows[(10, CEDAR)]["vol"] = aug
    rows[(11, CEDAR)]["vol"] = cents(q3 - jul - aug)

    # Northgate: exactly 28.4% of TTM debit volume
    others = sum(r["vol"] for (i, mid), r in rows.items() if mid != NORTHGATE)
    ng_total = cents(others * NORTHGATE_SHARE / (1 - NORTHGATE_SHARE))
    w = [SEASON[BNPL][i] * (1.012 ** i) * rng.uniform(0.98, 1.02) for i in range(12)]
    ws = sum(w)
    alloc = [cents(ng_total * x / ws) for x in w]
    alloc[11] = cents(ng_total - sum(alloc[:11]))
    for i in range(12):
        rows[(i, NORTHGATE)]["vol"] = alloc[i]

    # ---- counts and returns
    for m in merchants:
        for i in range(12):
            if not active(m, i):
                continue
            r = rows[(i, m["id"])]
            d = int(round(r["vol"] / m["ticket"] * rng.uniform(0.97, 1.03)))
            if m["id"] == SUMMIT and i in SUMMIT_RATES:
                d, u = solve_rate(rng, SUMMIT_RATES[i], d)
            else:
                base = m["unauth"]
                if m["id"] == SUMMIT:
                    base = [0.19, 0.21, 0.20, 0.22, 0.23, 0.24, 0.26, 0.27][i]
                u = int(round(d * base / 100.0 * rng.uniform(0.85, 1.15)))
                u = min(u, int(d * 0.0032))
            r["count"] = d
            r["unauth"] = u
            drift = ADMIN_DRIFT[0] / ADMIN_DRIFT[1] + (1 - ADMIN_DRIFT[0] / ADMIN_DRIFT[1]) * i / 11.0
            r["admin_raw"] = d * m["admin"] / 100.0 * drift * rng.uniform(0.95, 1.05)
            r["nsf"] = int(round(d * m["nsf"] / 100.0 * rng.uniform(0.9, 1.1)))
            r["credit_count"] = int(round(d * m["credit_ratio"] * rng.uniform(0.85, 1.15)))
            r["credit_vol"] = cents(r["credit_count"] * m["credit_ticket"] * rng.uniform(0.9, 1.1))
            r["card_vol"] = cents(r["vol"] * m["card_share"] * rng.uniform(0.9, 1.1)) if m["card_share"] else 0.0

    # Administrative returns: scale each month so the portfolio rate drifts 1.9% -> 2.6%
    for i in range(12):
        t = ADMIN_DRIFT[0] + (ADMIN_DRIFT[1] - ADMIN_DRIFT[0]) * i / 11.0
        if 0 < i < 11:
            t += rng.uniform(-0.05, 0.05)
        ids = [mid for (j, mid) in rows if j == i]
        total_d = sum(rows[(i, mid)]["count"] for mid in ids)
        raw = sum(rows[(i, mid)]["admin_raw"] for mid in ids)
        scale = (t / 100.0 * total_d) / raw
        for mid in ids:
            rows[(i, mid)]["admin"] = int(round(rows[(i, mid)]["admin_raw"] * scale))
        # nudge the largest merchant so the endpoints round exactly
        if i in (0, 11):
            tot = sum(rows[(i, mid)]["admin"] for mid in ids)
            want = int(round(t / 100.0 * total_d))
            rows[(i, NORTHGATE)]["admin"] += want - tot

    for key, r in rows.items():
        three = r["unauth"] + r["admin"] + r["nsf"]
        r["total"] = three + max(1, int(round(three * rng.uniform(0.025, 0.055))))
        del r["admin_raw"]

    # ---- revenue at the tier in the master
    def revenue(m, r, tier):
        fee, bps = TIERS[tier]
        return r["count"] * fee + r["vol"] * bps / 10000.0 + r["credit_count"] * CREDIT_FEE + r["card_vol"] * CARD_BPS / 10000.0

    for (i, mid), r in rows.items():
        r["revenue"] = revenue(byid[mid], r, byid[mid]["tier"])

    # ---- planted figures
    ttm_vol = cents(sum(r["vol"] for r in rows.values()))
    ttm_rev = cents(sum(r["revenue"] for r in rows.values()))
    per = {}
    for (i, mid), r in rows.items():
        p = per.setdefault(mid, {"vol": 0.0, "rev": 0.0})
        p["vol"] += r["vol"]
        p["rev"] += r["revenue"]
    ranked = sorted(per.items(), key=lambda kv: -kv[1]["vol"])
    ng_share = round(100.0 * per[NORTHGATE]["vol"] / ttm_vol, 1)
    second = ranked[1]
    harrow_under = cents(sum(rows[(i, HARROW)]["count"] * (TIERS["C"][0] - TIERS["B"][0])
                             + rows[(i, HARROW)]["vol"] * (TIERS["C"][1] - TIERS["B"][1]) / 10000.0 for i in range(6, 12)))
    cedar_q2 = cents(sum(rows[(i, CEDAR)]["vol"] for i in Q2))
    cedar_q3 = cents(sum(rows[(i, CEDAR)]["vol"] for i in Q3))
    cedar_pct = round(100.0 * (cedar_q3 / cedar_q2 - 1))
    summit = {i: pct(rows[(i, SUMMIT)]["unauth"], rows[(i, SUMMIT)]["count"]) for i in range(12)}
    summit_admin = {i: pct(rows[(i, SUMMIT)]["admin"], rows[(i, SUMMIT)]["count"]) for i in range(12)}
    summit_overall = {i: pct(rows[(i, SUMMIT)]["total"], rows[(i, SUMMIT)]["count"]) for i in range(12)}

    def admin_rate(i):
        ids = [mid for (j, mid) in rows if j == i]
        return round(100.0 * sum(rows[(i, mid)]["admin"] for mid in ids) / sum(rows[(i, mid)]["count"] for mid in ids), 1)

    def unauth_rate(i):
        ids = [mid for (j, mid) in rows if j == i]
        return round(100.0 * sum(rows[(i, mid)]["unauth"] for mid in ids) / sum(rows[(i, mid)]["count"] for mid in ids), 2)

    may = sum(r["vol"] for (i, mid), r in rows.items() if i == 7)
    return {
        "merchants": merchants, "byid": byid, "rows": rows,
        "ttm_vol": ttm_vol, "ttm_rev": ttm_rev, "per": per, "ranked": ranked,
        "ng_share": ng_share, "ng_rev": cents(per[NORTHGATE]["rev"]), "ng_aug_vol": rows[(10, NORTHGATE)]["vol"],
        "second": (second[0], round(100.0 * second[1]["vol"] / ttm_vol, 1)),
        "harrow_under": harrow_under, "cedar_q2": cedar_q2, "cedar_q3": cedar_q3, "cedar_pct": cedar_pct,
        "summit": summit, "summit_admin": summit_admin, "summit_overall": summit_overall,
        "admin_oct": admin_rate(0), "admin_sep": admin_rate(11),
        "unauth_sep": unauth_rate(11),
        "may_inflation": round(100.0 * may / ttm_vol, 1),
        "blended_bps": round(ttm_rev / ttm_vol * 10000.0, 1),
    }


# ----------------------------------------------------------------- writers
def write_activity_csv(path, model, i):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for m in model["merchants"]:
            key = (i, m["id"])
            if key not in model["rows"]:
                continue
            r = model["rows"][key]
            w.writerow([month_label(i), m["id"], m["name"], m["vertical"], r["count"], "%.2f" % r["vol"],
                        r["credit_count"], "%.2f" % r["credit_vol"], r["total"], r["unauth"], r["admin"], r["nsf"],
                        "%.2f" % r["card_vol"]])


def write_master_xlsx(path, model):
    wb = Workbook()
    wb.properties.creator = "Ironbridge"
    wb.properties.lastModifiedBy = "Ironbridge"
    wb.properties.title = ""
    ws = wb.active
    ws.title = "Merchants"
    ws.append(["merchant_id", "merchant_name", "vertical", "onboarded", "account_manager", "pricing_tier", "contract_renewal", "status"])
    for m in model["merchants"]:
        ws.append([m["id"], m["name"], m["vertical"], date.fromisoformat(m["onboarded"]), m["am"], m["tier"],
                   date.fromisoformat(m["renewal"]), m["status"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7E6E6")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"
        row[6].number_format = "yyyy-mm-dd"
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(40, max(12, width + 2))
    wb.save(path)


def new_pdf():
    pdf = FPDF(format="letter")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_creation_date(datetime(2026, 1, 6, 10, 2, 0))
    pdf.set_margins(20, 20, 20)
    pdf.add_page()
    return pdf


def write_pricing_pdf(path):
    pdf = new_pdf()
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "IRONBRIDGE PAYMENTS", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.cell(0, 5, "1400 Patterson Boulevard, Suite 300, Dayton, OH 45402", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(8)
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Merchant Pricing Schedule 2026", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 6, "Effective January 1, 2026. Supersedes the 2025 schedule.", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7, "ACH debit pricing by tier", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    widths = [18, 32, 58, 64]
    pdf.set_fill_color(230, 230, 230)
    pdf.set_font("Helvetica", "B", 10)
    for wdt, h in zip(widths, ["Tier", "Per debit item", "Basis points on debit volume", "Monthly debit volume"]):
        pdf.cell(wdt, 8, h, border=1, fill=True)
    pdf.ln(8)
    pdf.set_font("Helvetica", size=10)
    for row in content.PRICING_ROWS:
        for wdt, c in zip(widths, row):
            pdf.cell(wdt, 8, c, border=1)
        pdf.ln(8)
    pdf.ln(5)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7, "Other items", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    for line in content.PRICING_OTHER:
        pdf.multi_cell(0, 5.5, line, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7, "How monthly fees are calculated", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(0, 5.5, content.PRICING_FORMULA, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7, "Tier assignment", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(0, 5.5, content.PRICING_TIERING, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(8)
    pdf.set_font("Helvetica", "I", 9)
    pdf.cell(0, 5, "Approved: Helen Varga, Chief Risk Officer, and Greg Tanaka, Director of Merchant Contracts. December 12, 2025.", new_x="LMARGIN", new_y="NEXT")
    pdf.output(path)


def write_memo_docx(path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    doc.add_heading("Ironbridge Payments", level=0)
    doc.add_heading("Merchant Risk Policy", level=1)
    for k, v in content.MEMO_HEADER:
        p = doc.add_paragraph()
        run = p.add_run(k + "  ")
        run.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    for section in content.MEMO_SECTIONS:
        doc.add_heading(section["title"], level=2)
        for para in section.get("paras", []):
            doc.add_paragraph(para)
        if "table" in section:
            hdr, body = section["table"]
            t = doc.add_table(rows=1, cols=len(hdr))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(hdr):
                cell = t.rows[0].cells[j]
                cell.text = h
                cell.paragraphs[0].runs[0].bold = True
            for row in body:
                cells = t.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = v
            doc.add_paragraph()
        for para in section.get("after", []):
            doc.add_paragraph(para)
    p = doc.add_paragraph(content.MEMO_SIGNOFF)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cp = doc.core_properties
    cp.author = "Ironbridge"
    cp.last_modified_by = "Ironbridge"
    cp.created = datetime(2026, 1, 12, 9, 30)
    cp.modified = datetime(2026, 1, 12, 9, 30)
    doc.save(path)


def write_email(path, msg):
    with open(path, "w") as f:
        f.write("From: %s\n" % msg["from"])
        f.write("To: %s\n" % msg["to"])
        if msg.get("cc"):
            f.write("Cc: %s\n" % msg["cc"])
        f.write("Date: %s\n" % msg["date"])
        f.write("Subject: %s\n\n" % msg["subject"])
        f.write(msg["body"].rstrip("\n") + "\n")


KEEP = {"README.md", "brand-guidelines.md"}


def clean_dir(d):
    if os.path.isdir(d):
        for name in os.listdir(d):
            if name in KEEP:
                continue
            p = os.path.join(d, name)
            if os.path.isdir(p):
                shutil.rmtree(p)
            else:
                os.remove(p)
    os.makedirs(d, exist_ok=True)


def main():
    print("Ironbridge workshop pack generator (seed %d)" % SEED)
    model = build_model()
    clean_dir(D01)

    for i, fn in ACTIVITY_FILES:
        if fn == DUPLICATE_FILE[1]:
            continue
        write_activity_csv(os.path.join(D01, fn), model, i)
    shutil.copyfile(os.path.join(D01, DUPLICATE_FILE[0]), os.path.join(D01, DUPLICATE_FILE[1]))
    write_master_xlsx(os.path.join(D01, "merchant master.xlsx"), model)
    write_pricing_pdf(os.path.join(D01, "Pricing Schedule 2026.pdf"))
    write_memo_docx(os.path.join(D01, "Risk Policy memo.docx"))
    blank = new_pdf()
    blank.output(os.path.join(D01, BLANK_SCAN))
    for msg in content.emails(model):
        write_email(os.path.join(D01, msg["filename"]), msg)
    with open(os.path.join(D01, "lunch order thursday.txt"), "w") as f:
        f.write(content.LUNCH_ORDER)
    with open(os.path.join(D01, "teams-chat-merchant-ops-sept.txt"), "w") as f:
        f.write(content.teams_chat(model))

    n = len([f for f in os.listdir(D01) if f not in KEEP])
    print("  01-portfolio-inbox: %d files (13 activity exports incl. 1 duplicate, master xlsx, pricing PDF, policy docx, 4 emails, 1 chat export, 1 blank scan)" % n)
    rows = model["rows"]
    print("  TTM debit volume (May once): $%s  |  with May twice: $%s (+%.1f%%)" % (
        f"{model['ttm_vol']:,.2f}", f"{model['ttm_vol'] + sum(r['vol'] for (i, mid), r in rows.items() if i == 7):,.2f}", model["may_inflation"]))
    print("  TTM revenue at master tiers: $%s  (blended %.1f bps all-in)" % (f"{model['ttm_rev']:,.2f}", model["blended_bps"]))
    print("  Portfolio unauthorized rate, Sep 2026: %.2f%%" % model["unauth_sep"])
    print("  PLANTED  Summit Ridge unauthorized rate: Jun %.2f%%  Jul %.2f%%  Aug %.2f%%  Sep %.2f%%  (admin Sep %.2f%%, overall Sep %.2f%%)" % (
        model["summit"][8], model["summit"][9], model["summit"][10], model["summit"][11], model["summit_admin"][11], model["summit_overall"][11]))
    print("           Summit Ridge Oct-May: " + "  ".join("%.2f" % model["summit"][i] for i in range(8)))
    print("  PLANTED  Northgate BNPL: %.1f%% of TTM debit volume ($%s); TTM revenue $%s (largest line); next largest %s at %.1f%%" % (
        model["ng_share"], f"{model['per'][NORTHGATE]['vol']:,.2f}", f"{model['ng_rev']:,.2f}", model["byid"][model["second"][0]]["name"], model["second"][1]))
    print("  PLANTED  Harrow & Finch under-billed Apr-Sep 2026 (billed Tier B, amendment says Tier C): $%s" % f"{model['harrow_under']:,.2f}")
    print("  PLANTED  Cedar Falls Q2 $%s -> Q3 $%s (%d%%); Jul/Aug/Sep $%s / $%s / $%s" % (
        f"{model['cedar_q2']:,.2f}", f"{model['cedar_q3']:,.2f}", model["cedar_pct"],
        f"{rows[(9, CEDAR)]['vol']:,.2f}", f"{rows[(10, CEDAR)]['vol']:,.2f}", f"{rows[(11, CEDAR)]['vol']:,.2f}"))
    print("  PLANTED  duplicate: '%s' is byte-identical to '%s'" % (DUPLICATE_FILE[1], DUPLICATE_FILE[0]))
    print("  PLANTED  portfolio administrative return rate: %.1f%% (Oct 2025) -> %.1f%% (Sep 2026)" % (model["admin_oct"], model["admin_sep"]))
    print("  PLANTED  needs a human: %s (blank page); 'lunch order thursday.txt' (not work, undated)" % BLANK_SCAN)
    print("done")


if __name__ == "__main__":
    main()
